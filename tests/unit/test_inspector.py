"""ContentExtractor 单元测试。"""

from __future__ import annotations

import tempfile
from pathlib import Path

from app.core.inspector import (
    MAX_CHARS,
    MAX_SIZE_BYTES,
    SUPPORTED_EXTENSIONS,
    ExtractResult,
    _should_extract,
    extract_for_entries,
)
from app.core.models import FileEntry


def _make_entry(
    tmp: Path,
    name: str,
    size: int = 100,
    content: bytes = b"",
) -> FileEntry:
    path = tmp / name
    path.write_bytes(content)
    return FileEntry(
        path=path,
        name=name,
        ext=name[name.rfind(".") :].lower() if "." in name else "",
        size=size,
        mtime=0.0,
        is_hidden=False,
        depth=1,
    )


class TestShouldExtract:
    def test_supported_ext_allowed(self, tmp_path: Path) -> None:
        entry = _make_entry(tmp_path, "a.pdf", size=100)
        assert _should_extract(entry) is True

    def test_unsupported_ext_rejected(self, tmp_path: Path) -> None:
        entry = _make_entry(tmp_path, "a.exe", size=100)
        assert _should_extract(entry) is False

    def test_oversized_rejected(self, tmp_path: Path) -> None:
        entry = _make_entry(tmp_path, "a.pdf", size=MAX_SIZE_BYTES + 1)
        assert _should_extract(entry) is False

    def test_at_boundary_rejected(self, tmp_path: Path) -> None:
        entry = _make_entry(tmp_path, "a.pdf", size=MAX_SIZE_BYTES)
        assert _should_extract(entry) is False

    def test_all_supported_extensions(self, tmp_path: Path) -> None:
        for ext in SUPPORTED_EXTENSIONS:
            entry = _make_entry(tmp_path, f"file{ext}", size=100)
            assert _should_extract(entry) is True


class TestExtractText:
    def test_extract_utf8(self, tmp_path: Path) -> None:
        text = "这是一段中文测试文本，用于验证提取功能。" * 10
        entry = _make_entry(tmp_path, "test.txt", size=len(text.encode("utf-8")), content=text.encode("utf-8"))
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        assert "中文测试" in results[0].text_head

    def test_extract_gbk(self, tmp_path: Path) -> None:
        text = "GBK编码测试内容" * 20
        entry = _make_entry(tmp_path, "test.txt", size=len(text.encode("gbk")), content=text.encode("gbk"))
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        # chardet 应该能探测到 GBK 编码
        assert "GBK" in results[0].text_head or "编码" in results[0].text_head

    def test_extract_md(self, tmp_path: Path) -> None:
        text = "# Markdown 标题\n\n正文内容。" * 10
        entry = _make_entry(tmp_path, "readme.md", size=len(text.encode("utf-8")), content=text.encode("utf-8"))
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        assert "Markdown" in results[0].text_head

    def test_truncated_to_max_chars(self, tmp_path: Path) -> None:
        text = "A" * (MAX_CHARS * 2)
        entry = _make_entry(tmp_path, "big.txt", size=len(text), content=text.encode("utf-8"))
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert len(results[0].text_head) == MAX_CHARS


class TestExtractDocx:
    def test_extract_docx(self, tmp_path: Path) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("这是 Word 文档的正文内容。")
        doc.add_paragraph("第二段内容。")
        path = tmp_path / "test.docx"
        doc.save(str(path))

        entry = FileEntry(
            path=path,
            name="test.docx",
            ext=".docx",
            size=path.stat().st_size,
            mtime=0.0,
            is_hidden=False,
            depth=1,
        )
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        assert "Word" in results[0].text_head


class TestExtractXlsx:
    def test_extract_xlsx(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "表格内容"
        ws["B1"] = "测试数据"
        path = tmp_path / "test.xlsx"
        wb.save(str(path))

        entry = FileEntry(
            path=path,
            name="test.xlsx",
            ext=".xlsx",
            size=path.stat().st_size,
            mtime=0.0,
            is_hidden=False,
            depth=1,
        )
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        assert "表格内容" in results[0].text_head


class TestExtractPptx:
    def test_extract_pptx(self, tmp_path: Path) -> None:
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.placeholders[0].text = "演示文稿标题"
        path = tmp_path / "test.pptx"
        prs.save(str(path))

        entry = FileEntry(
            path=path,
            name="test.pptx",
            ext=".pptx",
            size=path.stat().st_size,
            mtime=0.0,
            is_hidden=False,
            depth=1,
        )
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].text_head is not None
        assert "演示文稿" in results[0].text_head


class TestExtractPdf:
    def test_extract_pdf_unsupported_without_pypdf(self, tmp_path: Path) -> None:
        """如果 pypdf 不可用，应返回 error 而非崩溃。"""
        entry = _make_entry(tmp_path, "test.pdf", size=100, content=b"%PDF-1.4 fake")
        results = extract_for_entries([entry])
        # pypdf 应该能装上，这里验证不会崩溃
        assert len(results) == 1


class TestErrorHandling:
    def test_corrupt_file_returns_error(self, tmp_path: Path) -> None:
        """损坏的文件应返回 error 而非抛异常。"""
        entry = _make_entry(tmp_path, "corrupt.docx", size=100, content=b"not a real docx")
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].error is not None

    def test_nonexistent_file_returns_error(self, tmp_path: Path) -> None:
        """不存在的文件应返回 error。"""
        entry = FileEntry(
            path=tmp_path / "nonexistent.txt",
            name="nonexistent.txt",
            ext=".txt",
            size=0,
            mtime=0.0,
            is_hidden=False,
            depth=1,
        )
        results = extract_for_entries([entry])
        assert len(results) == 1
        assert results[0].error is not None


class TestBatchExtract:
    def test_mixed_entries(self, tmp_path: Path) -> None:
        """混合支持与不支持的扩展名。"""
        text_entry = _make_entry(tmp_path, "a.txt", size=10, content=b"hello")
        exe_entry = _make_entry(tmp_path, "b.exe", size=10, content=b"MZ")
        results = extract_for_entries([text_entry, exe_entry])
        # 只有 txt 被提取
        assert len(results) == 1
        assert results[0].path.name == "a.txt"

    def test_empty_list(self, tmp_path: Path) -> None:
        results = extract_for_entries([])
        assert results == []
